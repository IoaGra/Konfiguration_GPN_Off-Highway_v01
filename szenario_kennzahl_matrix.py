"""Einlesen und Auswerten der Szenario-Kennzahl-Matrix."""

from __future__ import annotations

import os
from openpyxl import load_workbook


def lade_matrix(pfad: str = "szenario_kennzahl_matrix.xlsx"):
    if not os.path.exists(pfad):
        raise FileNotFoundError(f"Die Szenario-Kennzahl-Matrix fehlt: {pfad}")

    arbeitsmappe = load_workbook(pfad, data_only=True)
    arbeitsblatt = arbeitsmappe.active
    zeilen = list(arbeitsblatt.iter_rows(values_only=True))
    if not zeilen:
        raise ValueError("Die Szenario-Kennzahl-Matrix enthält keine Daten.")

    kopf = [str(z).strip() if z is not None else "" for z in zeilen[0]]
    szenario_typen = [kopf[i] for i in range(2, len(kopf)) if kopf[i]]

    matrix = {}
    for zeile in zeilen[1:]:
        if not zeile or len(zeile) < 2 or zeile[1] is None:
            continue
        kennzahl = str(zeile[1]).strip()
        matrix[kennzahl] = {}
        for index, typ in enumerate(szenario_typen):
            zelle = zeile[2 + index] if 2 + index < len(zeile) else 0
            try:
                matrix[kennzahl][typ] = 1 if int(zelle or 0) == 1 else 0
            except (TypeError, ValueError):
                matrix[kennzahl][typ] = 0
    return szenario_typen, matrix


def aktive_kennzahlen(matrix: dict, szenario_typ: str) -> list[str]:
    return [
        kennzahl
        for kennzahl, zuordnung in matrix.items()
        if zuordnung.get(szenario_typ, 0) == 1
    ]


def vereinigte_kennzahlen(matrix: dict, szenario_typen: list[str]) -> list[str]:
    typen = set(szenario_typen)
    return [
        kennzahl
        for kennzahl, zuordnung in matrix.items()
        if any(zuordnung.get(typ, 0) == 1 for typ in typen)
    ]
