"""Einlesen der quantitativen Datengrundlage aus ``datengrundlage.xlsx``."""

from __future__ import annotations

import os
from openpyxl import load_workbook

FEHLEND_MARKER = {"/", "unbekannt", "", "none", "null", "nan"}


def _ist_fehlend(zelle) -> bool:
    if zelle is None:
        return True
    if isinstance(zelle, str):
        return zelle.strip().lower() in FEHLEND_MARKER
    return False


def lade_daten(pfad: str = "datengrundlage.xlsx"):
    """Liest Kennzahlen, Referenzwerte, Optionen und vorhandene Istwerte ein."""
    if not os.path.exists(pfad):
        raise FileNotFoundError(f"Die Datengrundlage wurde nicht gefunden: {pfad}")

    arbeitsmappe = load_workbook(pfad, data_only=True)
    arbeitsblatt = arbeitsmappe.active
    zeilen = list(arbeitsblatt.iter_rows(values_only=True))
    if not zeilen:
        raise ValueError("Die Datengrundlage enthält keine Daten.")

    kopf = [str(z).strip() if z is not None else "" for z in zeilen[0]]
    optionen = [kopf[i] for i in range(6, len(kopf)) if kopf[i]]

    kpis = {}
    werte = {option: {} for option in optionen}
    fehlend = {option: [] for option in optionen}

    for zeile in zeilen[1:]:
        if not zeile or len(zeile) < 6 or zeile[1] is None:
            continue
        kennzahl = str(zeile[1]).strip()
        if not kennzahl or kennzahl.lower().startswith("hinweis"):
            continue

        try:
            zielwert = float(zeile[4])
            grenzwert = float(zeile[5])
        except (TypeError, ValueError) as fehler:
            raise ValueError(
                f"Ziel- oder Grenzwert der Kennzahl '{kennzahl}' ist ungültig."
            ) from fehler

        kpis[kennzahl] = {
            "dimension": str(zeile[0]).strip() if zeile[0] is not None else "",
            "formelzeichen": str(zeile[2]).strip() if zeile[2] is not None else "",
            "einheit": str(zeile[3]).strip() if zeile[3] is not None else "",
            "zielwert": zielwert,
            "grenzwert": grenzwert,
        }

        for index, option in enumerate(optionen):
            zelle = zeile[6 + index] if 6 + index < len(zeile) else None
            if _ist_fehlend(zelle):
                fehlend[option].append(kennzahl)
                continue
            try:
                werte[option][kennzahl] = float(zelle)
            except (TypeError, ValueError):
                fehlend[option].append(kennzahl)

    return kpis, optionen, werte, fehlend


def dimensionen_aus_kpis(kpis: dict) -> list[str]:
    """Gibt die Leistungsdimensionen in stabiler Reihenfolge zurück."""
    dimensionen: list[str] = []
    for info in kpis.values():
        dimension = info.get("dimension", "")
        if dimension and dimension not in dimensionen:
            dimensionen.append(dimension)
    return dimensionen
