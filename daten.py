"""
daten.py
========
Liest die Datei 'daten.xlsx' mit der Datengrundlage:

  Leistungsdimension | Kennzahl | Formelzeichen | Einheit | Zielwert | Grenzwert | Standort A | Standort B | ...

Fehlende Werte können in der Excel-Datei mit '/', 'unbekannt' oder leeren
Zellen markiert werden. Diese Werte werden später im Dialog abgefragt.
"""

import os
from openpyxl import load_workbook

FEHLEND_MARKER = {"/", "unbekannt", "", "none", "null", "nan"}


def _ist_fehlend(zelle) -> bool:
    if zelle is None:
        return True
    if isinstance(zelle, str):
        return zelle.strip().lower() in FEHLEND_MARKER
    return False


def lade_daten(pfad: str = "daten.xlsx"):
    """Liest daten.xlsx und gibt zurück:

    kpis      : {kennzahl: {dimension, formelzeichen, einheit, zielwert, grenzwert}}
    optionen  : Liste der Options-Spalten
    werte     : {option: {kennzahl: float}}   nur vorhandene Werte
    fehlend   : {option: [kennzahl, ...]}     fehlende Werte
    """
    if not os.path.exists(pfad):
        return {}, [], {}, {}

    wb = load_workbook(pfad, data_only=True)
    ws = wb.active
    zeilen = list(ws.iter_rows(values_only=True))
    if not zeilen:
        return {}, [], {}, {}

    kopf = [str(z).strip() if z is not None else "" for z in zeilen[0]]
    optionen = [kopf[i] for i in range(6, len(kopf)) if kopf[i]]

    kpis = {}
    werte = {o: {} for o in optionen}
    fehlend = {o: [] for o in optionen}

    for z in zeilen[1:]:
        if not z or z[1] is None:
            continue
        kennzahl = str(z[1]).strip()
        if not kennzahl or kennzahl.lower().startswith("hinweis"):
            continue

        kpis[kennzahl] = {
            "dimension": str(z[0]).strip() if z[0] is not None else "",
            "formelzeichen": str(z[2]).strip() if z[2] is not None else "",
            "einheit": str(z[3]).strip() if z[3] is not None else "",
            "zielwert": float(z[4]),
            "grenzwert": float(z[5]),
        }

        for i, opt in enumerate(optionen):
            zelle = z[6 + i] if 6 + i < len(z) else None
            if _ist_fehlend(zelle):
                fehlend[opt].append(kennzahl)
            else:
                try:
                    werte[opt][kennzahl] = float(zelle)
                except (TypeError, ValueError):
                    fehlend[opt].append(kennzahl)

    return kpis, optionen, werte, fehlend


def dimensionen_aus_kpis(kpis: dict) -> list:
    """Dimensionen in stabiler Reihenfolge des ersten Auftretens."""
    gesehen = []
    for info in kpis.values():
        dim = info.get("dimension", "")
        if dim and dim not in gesehen:
            gesehen.append(dim)
    return gesehen
