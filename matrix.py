"""
matrix.py  (V6)
===============
Liest die Szenario-Kennzahl-Matrix aus 'szenario_kennzahl_matrix.xlsx'.

  Leistungsdimension | Kennzahl | <Szenario-Typ 1> | <Szenario-Typ 2> | ...

1 = Kennzahl ist fuer diesen Szenario-Typ relevant, 0 = nicht relevant.
Der Kennzahl-Name ist der Schluessel (identisch zu daten.xlsx).
"""

import os
from openpyxl import load_workbook


def lade_matrix(pfad: str = "szenario_kennzahl_matrix.xlsx"):
    """Gibt zurueck:
    szenario_typen : Liste der Typ-Spalten (lesbare Namen aus der Datei)
    matrix         : {kennzahl: {typ: 0|1}}
    """
    if not os.path.exists(pfad):
        return [], {}

    wb = load_workbook(pfad, data_only=True)
    ws = wb.active
    zeilen = list(ws.iter_rows(values_only=True))
    if not zeilen:
        return [], {}

    kopf = [str(z).strip() if z is not None else "" for z in zeilen[0]]
    szenario_typen = [kopf[i] for i in range(2, len(kopf)) if kopf[i]]

    matrix = {}
    for z in zeilen[1:]:
        if not z or z[1] is None:
            continue
        kennzahl = str(z[1]).strip()
        matrix[kennzahl] = {}
        for i, typ in enumerate(szenario_typen):
            zelle = z[2 + i] if 2 + i < len(z) else None
            matrix[kennzahl][typ] = int(zelle) if zelle is not None else 0
    return szenario_typen, matrix


def aktive_kennzahlen(matrix: dict, typ: str) -> list:
    """Kennzahlen, die fuer den gewaehlten Szenario-Typ aktiv (=1) sind."""
    return [k for k, spalten in matrix.items() if spalten.get(typ, 0) == 1]
